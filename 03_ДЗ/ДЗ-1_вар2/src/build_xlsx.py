"""Excel с расчётами ДЗ-1, вариант 2. Все величины — формулами."""
import os
import sys
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from calc import LAYERS, OTTM, BITS, FOREIGN, WALL_CHOICE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "ДЗ1_вар2_расчеты.xlsx")

wb = Workbook()
th = Side(style="thin", color="404040")
BOX = Border(left=th, right=th, top=th, bottom=th)
HEAD = PatternFill("solid", fgColor="D5E8E4")
SUBH = PatternFill("solid", fgColor="EEF5F3")
INP = PatternFill("solid", fgColor="FFF4D6")
RES = PatternFill("solid", fgColor="FCE4D6")
B = Font(bold=True)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
LW = Alignment(horizontal="left", vertical="center", wrap_text=True)


def fmt(ws, rng, fill=None, bold=False, align=C, nf=None):
    for row in ws[rng]:
        for c in row:
            c.border = BOX
            c.alignment = align
            if fill:
                c.fill = fill
            if bold:
                c.font = B
            if nf:
                c.number_format = nf


def head(ws, row, col, items, fill=HEAD):
    for j, h in enumerate(items):
        ws.cell(row, col + j, h)
    fmt(ws, f"{get_column_letter(col)}{row}:{get_column_letter(col + len(items) - 1)}{row}", fill, True)


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="1F4E4A")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, size=10, color="404040")


def amin(ws, cell, vals, crit, x):
    ws[cell] = ArrayFormula(cell, f"=MIN(IF({crit}>={x},{vals}))")


# ===================================================================== Исходные данные
ws = wb.active
ws.title = "Исходные данные"
widths(ws, [44, 14, 30])
title(ws, "ДЗ № 1. Проектирование конструкции скважины. Вариант 2",
      "Жёлтые ячейки — исходные данные; все остальные листы считаются от них")
head(ws, 4, 1, ["Параметр", "Значение", "Примечание"])
params = [
    ("Назначение скважины", "газовая", "вариант 2 (№ 2 по журналу)"),
    ("Глубина скважины H, м", 2400, ""),
    ("Диаметр эксплуатационной колонны, мм", 127.0, "задан"),
    ("Неустойчивые породы до глубины, м", 500, "склонны к поглощениям"),
    ("g, м/с²", 9.81, ""),
    ("ρв, кг/м³", 1000, "плотность воды"),
    ("Ks при z ≤ 1200 м", 1.1, "ПБ, п. 2.7.3.3"),
    ("Ks при z > 1200 м", 1.05, "ПБ, п. 2.7.3.3"),
    ("Граница смены Ks, м", 1200, ""),
    ("Подъём цемента над башмаком пред. ОК, м", 500, "газовая скв., ПБ п. 2.7.4.11"),
    ("Зазор δ долото – внутр. стенка ОК, мм", 5, "принят из диапазона 5–10 мм"),
]
P = {}
for i, (n, v, c) in enumerate(params):
    r = 5 + i
    ws.cell(r, 1, n); ws.cell(r, 2, v); ws.cell(r, 3, c)
    P[n] = f"'Исходные данные'!$B${r}"
last = 4 + len(params)
fmt(ws, f"A5:C{last}", align=LW)
fmt(ws, f"B5:B{last}", INP)
r0 = last + 2
ws.cell(r0, 1, "Геологический разрез").font = B
head(ws, r0 + 1, 1, ["Порода", "Кровля, м", "Подошва, м", "Pпл, МПа", "Pпогл, МПа"])
ws.column_dimensions["D"].width = 12
ws.column_dimensions["E"].width = 12
LROW = []
for i, (t, b, rock, pf, pfr) in enumerate(LAYERS):
    r = r0 + 2 + i
    ws.cell(r, 1, rock); ws.cell(r, 2, t); ws.cell(r, 3, b); ws.cell(r, 4, pf); ws.cell(r, 5, pfr)
    LROW.append(r)
fmt(ws, f"A{LROW[0]}:E{LROW[-1]}", INP)
fmt(ws, f"A{LROW[0]}:A{LROW[-1]}", INP, align=LW)
ws.cell(LROW[-1] + 1, 1, "Давления отнесены к подошве интервала.").font = Font(italic=True, size=9)
G, RW = P["g, м/с²"], P["ρв, кг/м³"]
KS1, KS2, ZKS = P["Ks при z ≤ 1200 м"], P["Ks при z > 1200 м"], P["Граница смены Ks, м"]
CEMG = P["Подъём цемента над башмаком пред. ОК, м"]
DS = P["Зазор δ долото – внутр. стенка ОК, мм"]
HW, DP, UNS = P["Глубина скважины H, м"], P["Диаметр эксплуатационной колонны, мм"], P["Неустойчивые породы до глубины, м"]
SRC = "'Исходные данные'!"

# ===================================================================== Справочник (нужен для ссылок)
wsR = wb.create_sheet("Справочник ГОСТ")
widths(wsR, [11, 10, 9, 10, 12, 3, 14, 3, 12, 9])
title(wsR, "Справочные данные", "ГОСТ 632-80 (трубы ОТТМ), ГОСТ 20692-2003 (долота), зазоры Δ (семинар 1)")
head(wsR, 4, 1, ["Ø усл.", "D, мм", "s, мм", "d, мм", "Dн муфты"])
r = 5
for D in sorted(OTTM):
    nom, dm, walls = OTTM[D]
    for s, d in walls:
        wsR.cell(r, 1, nom); wsR.cell(r, 2, D); wsR.cell(r, 3, s); wsR.cell(r, 4, d); wsR.cell(r, 5, dm)
        r += 1
CL = r - 1
fmt(wsR, f"A5:E{CL}", nf="0.0")
head(wsR, 4, 7, ["Долото, мм"])
for i, b in enumerate(BITS):
    wsR.cell(5 + i, 7, b)
BL = 4 + len(BITS)
fmt(wsR, f"G5:G{BL}", nf="0.0")
head(wsR, 4, 9, ["Ø усл. до", "Δ, мм"])
for i, (lim, dl) in enumerate([(127, 8), (146, 10), (245, 13), (299, 18), (426, 22)]):
    wsR.cell(5 + i, 9, lim); wsR.cell(5 + i, 10, dl)
fmt(wsR, "I5:J9")
R = "'Справочник ГОСТ'!"
CN, CD, CS, Cd, CM = (f"{R}$A$5:$A${CL}", f"{R}$B$5:$B${CL}", f"{R}$C$5:$C${CL}", f"{R}$D$5:$D${CL}", f"{R}$E$5:$E${CL}")
BB = f"{R}$G$5:$G${BL}"
DLIM, DVAL = f"{R}$I$5:$I$9", f"{R}$J$5:$J$9"

# ===================================================================== Давления
w2 = wb.create_sheet("Давления", 1)
widths(w2, [26, 10, 10, 10, 10, 10, 9, 10, 12, 14])
title(w2, "1. Индексы давлений и относительная плотность бурового раствора",
      "KF = Pпл/(ρв·g·z);  KFR = Pпогл/(ρв·g·z);  ρR = Ks·KF;  ρб.р. = ρR·ρв")
head(w2, 4, 1, ["Порода", "z₁, м", "z₂, м", "Pпл, МПа", "Pпогл, МПа", "KF", "KFR", "Ks", "ρR", "ρб.р., кг/м³"])
subs = [(0, 0, None), (1, 1, None), (2, 2, "1200"), (3, 2, "1200b"), (4, 3, None)]
rows_sub = []
for i, (k, li, cut) in enumerate(subs):
    r = 5 + i
    lr = LROW[li]
    w2.cell(r, 1, f"={SRC}A{lr}")
    w2.cell(r, 2, f"={SRC}B{lr}" if cut != "1200b" else f"={ZKS}")
    w2.cell(r, 3, f"={SRC}C{lr}" if cut != "1200" else f"={ZKS}")
    w2.cell(r, 4, f"={SRC}D{lr}")
    w2.cell(r, 5, f"={SRC}E{lr}")
    w2.cell(r, 6, f"=D{r}*10^6/({RW}*{G}*{SRC}C{lr})")
    w2.cell(r, 7, f"=E{r}*10^6/({RW}*{G}*{SRC}C{lr})")
    w2.cell(r, 8, f"=IF(C{r}<={ZKS},{KS1},{KS2})")
    w2.cell(r, 9, f"=H{r}*F{r}")
    w2.cell(r, 10, f"=I{r}*{RW}")
    rows_sub.append(r)
fmt(w2, "A5:J9")
fmt(w2, "A5:A9", align=LW)
fmt(w2, "F5:I9", RES, nf="0.000")
fmt(w2, "J5:J9", RES, nf="0")
w2["A10"] = "Известняк (1100–2100 м) разделён на 1100–1200 и 1200–2100 м из-за смены Ks; KF и KFR в пласте постоянны."
w2["A10"].font = Font(italic=True, size=9)

w2["A12"] = "2. Совместимость условий бурения (весь открытый ствол — одна плотность раствора)"
w2["A12"].font = Font(bold=True, size=12, color="1F4E4A")
head(w2, 13, 1, ["Интервал бурения", "от, м", "до, м", "ρR прин.", "KF max", "KFR min", "Вывод"])
secs = [("под кондуктор", "0", "='Конструкция'!C6", "I5", "F5", "G5"),
        ("под промежуточную ОК", "='Конструкция'!C6", "='Конструкция'!D6", "MAX(I5:I6)", "MAX(F5:F6)", "MIN(G5:G6)"),
        ("под эксплуатационную ОК", "='Конструкция'!D6", "='Конструкция'!E6", "MAX(I7:I9)", "MAX(F7:F9)", "MIN(G7:G9)")]
for i, (n, a, b, rr, kf, kfr) in enumerate(secs):
    r = 14 + i
    w2.cell(r, 1, n); w2.cell(r, 2, f"={a}" if not a.startswith("=") else a); w2.cell(r, 3, b)
    w2.cell(r, 4, f"={rr}"); w2.cell(r, 5, f"={kf}"); w2.cell(r, 6, f"={kfr}")
    w2.cell(r, 7, f'=IF(AND(E{r}<D{r},D{r}<F{r}),"совместимы","несовместимы")')
fmt(w2, "A14:G16")
fmt(w2, "A14:A16", align=LW)
fmt(w2, "D14:F16", RES, nf="0.000")
w2["A18"] = "Можно ли вскрыть известняк, не перекрыв интервал 0–1100 м?"
w2["A18"].font = B
chk = [("ρR, нужная для 1100–2400 м", "=D16"),
       ("KFR min выше 1100 м", "=MIN(G5:G6)"),
       ("ρR, допустимая выше 1100 м", "=D15"),
       ("KF известняка", "=F7")]
for i, (n, f) in enumerate(chk):
    w2.cell(19 + i, 1, n); w2.cell(19 + i, 2, f)
fmt(w2, "A19:B22", align=LW)
fmt(w2, "B19:B22", RES, nf="0.000")
w2["A23"] = "Ответ"
w2["B23"] = ('=IF(OR(B19>=B20,B21<=B22),"нет — поглощение выше 1100 м / ГНВП в известняке; '
             'нужна промежуточная колонна до 1100 м","да")')
w2.merge_cells("B23:J23")
fmt(w2, "A23:J23", align=LW)
w2.row_dimensions[23].height = 30

# ===================================================================== График
w3 = wb.create_sheet("График", 2)
widths(w3, [9, 8, 9, 8, 9, 8, 9, 8])
title(w3, "Совмещённый график давлений", "Данные ступенчатых линий (ссылки на лист «Давления»)")
head(w3, 4, 1, ["KF", "z", "KFR", "z", "ρR", "z", "ρR прин.", "z"])
r = 5
for li, (srow_top, srow_bot) in enumerate([(5, 5), (6, 6), (7, 8), (9, 9)]):
    for z in (f"Давления!B{srow_top}", f"Давления!C{srow_bot}"):
        w3.cell(r, 1, f"=Давления!F{srow_top}"); w3.cell(r, 2, f"={z}")
        w3.cell(r, 3, f"=Давления!G{srow_top}"); w3.cell(r, 4, f"={z}")
        r += 1
kl = r - 1
r = 5
for sr in range(5, 10):
    for z in (f"Давления!B{sr}", f"Давления!C{sr}"):
        w3.cell(r, 5, f"=Давления!I{sr}"); w3.cell(r, 6, f"={z}")
        r += 1
rl = r - 1
r = 5
for sr in (14, 15, 16):
    for z in (f"Давления!B{sr}", f"Давления!C{sr}"):
        w3.cell(r, 7, f"=Давления!D{sr}"); w3.cell(r, 8, f"={z}")
        r += 1
pl = r - 1
fmt(w3, f"A5:H{rl}", nf="0.000")
ch = ScatterChart()
ch.title = "Совмещённый график давлений (вариант 2)"
ch.style = 2
ch.x_axis.title = "KF, KFR, ρR"
ch.y_axis.title = "Глубина, м"
ch.y_axis.scaling.orientation = "maxMin"
ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = 0, 2400, 300
ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = 0.7, 1.8, 0.1
ch.x_axis.number_format, ch.y_axis.number_format = "0.0", "0"
ch.x_axis.delete = ch.y_axis.delete = False
ch.height, ch.width = 17, 17
for xc, yc, lastr, name, color, dash, w in [
    (1, 2, kl, "KF", "000000", None, 28575), (3, 4, kl, "KFR", "000000", "dash", 28575),
    (5, 6, rl, "ρR = Ks·KF", "1F4E79", "sysDot", 19050), (7, 8, pl, "ρR принятая", "C00000", None, 15875)]:
    s = Series(Reference(w3, min_col=yc, min_row=5, max_row=lastr), Reference(w3, min_col=xc, min_row=5, max_row=lastr), title=name)
    s.marker.symbol = "none"
    s.smooth = False
    s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=color, w=w, prstDash=dash))
    ch.series.append(s)
w3.add_chart(ch, "J4")

# ===================================================================== Конструкция
w4 = wb.create_sheet("Конструкция", 3)
widths(w4, [50, 8, 15, 15, 17])
title(w4, "3. Конструкция скважины: колонны, цементирование, диаметры",
      "Диаметры — снизу вверх: DB = dм + 2Δ → долото ГОСТ 20692 ≥ DB;  d пред. ОК ≥ DB + 2δ → ОК ГОСТ 632-80")
head(w4, 5, 1, ["Показатель", "обозн.", "Кондуктор", "Промежуточная", "Эксплуатационная"])
labels = [
    ("Глубина спуска (башмак), м", "L"),                       # 6
    ("Кровля цемента (расчёт: башмак пред. ОК − 500 м), м", "H расч"),  # 7
    ("Кровля цемента принятая, м", "H"),                       # 8
    ("Требуемый внутренний диаметр d ≥ DB(ниже) + 2δ, мм", "d min"),  # 9
    ("Наружный диаметр ОК, мм", "D"),                          # 10
    ("Условный диаметр, мм", "Dусл"),                          # 11
    ("Толщина стенки (предварительно), мм", "s"),             # 12
    ("Внутренний диаметр, мм", "d"),                           # 13
    ("Диаметр муфты ОТТМ, мм", "dм"),                          # 14
    ("Радиальный зазор муфта–стенка, мм", "Δ"),                # 15
    ("Расчётный диаметр долота, мм", "DB расч"),               # 16
    ("Диаметр долота по ГОСТ 20692, мм", "DB"),                # 17
    ("Зазор долота ниже лежащего интервала в этой ОК, мм", "δ факт"),  # 18
]
for i, (a, b) in enumerate(labels):
    w4.cell(6 + i, 1, a); w4.cell(6 + i, 2, b)
w4["C6"] = f"={UNS}+10"
w4["D6"] = f"={SRC}C{LROW[1]}"
w4["E6"] = f"={HW}"
w4["C7"] = "до устья"
w4["D7"] = f"=C6-{CEMG}"
w4["E7"] = f"=D6-{CEMG}"
w4["C8"] = 0
w4["D8"] = "=IF(D7<50,0,D7)"
w4["E8"] = "=IF(E7<50,0,E7)"
# диаметры: эксплуатационная (E) -> промежуточная (D) -> кондуктор (C)
w4["E9"] = "—"
w4["D9"] = f"=E17+2*{DS}"
w4["C9"] = f"=D17+2*{DS}"
w4["E10"] = f"={DP}"
amin(w4, "D10", CD, Cd, "D9")
amin(w4, "C10", CD, Cd, "C9")
for col in "CDE":
    w4[f"{col}11"] = f"=INDEX({CN},MATCH({col}10,{CD},0))"
    w4[f"{col}14"] = f"=INDEX({CM},MATCH({col}10,{CD},0))"
    amin(w4, f"{col}15", DVAL, DLIM, f"{col}11")
    w4[f"{col}16"] = f"={col}14+2*{col}15"
    amin(w4, f"{col}17", BB, BB, f"{col}16")
w4["E12"] = "по расчёту на прочность"
w4["D12"] = WALL_CHOICE["Промежуточная"]
w4["C12"] = WALL_CHOICE["Кондуктор"]
w4["E13"] = "—"
w4["D13"] = f"=SUMIFS({Cd},{CD},D10,{CS},D12)"
w4["C13"] = f"=SUMIFS({Cd},{CD},C10,{CS},C12)"
w4["E18"] = "—"
w4["D18"] = "=(D13-E17)/2"
w4["C18"] = "=(C13-D17)/2"
fmt(w4, "A6:E18")
fmt(w4, "A6:A18", align=LW)
fmt(w4, "C9:E18", nf="0.0")
fmt(w4, "C12:D12", INP, nf="0.0")
for rr in (6, 8, 10, 17):
    fmt(w4, f"C{rr}:E{rr}", RES, True, nf="0.0" if rr in (10, 17) else "0")
w4["A20"] = "Обоснование"
w4["A20"].font = B
notes = [
    "Кондуктор: перекрывает неустойчивые породы, склонные к поглощениям (0–500 м); башмак на 10 м ниже, цемент до устья.",
    "Промежуточная: интервал 0–1100 м несовместим с известняком с АВПД (нужна ρR = 1,50 > KFR = 1,33 и 1,43 выше; "
    "ρR = 1,12 < KF = 1,33 известняка). Башмак — в подошве алевролитов. Расчётная кровля цемента 10 м → цементируем до устья.",
    "Эксплуатационная: до проектной глубины; кровля цемента 600 м (500 м над башмаком промежуточной).",
    "Кондуктор 273,1 мм: подходит только толщина 8,9 мм (d = 255,3 ≥ 254,5 мм); ОК 244,5 мм не проходит (d max = 228,7 мм).",
]
for i, t in enumerate(notes):
    c = w4.cell(21 + i, 1, f"{i + 1}) {t}")
    w4.merge_cells(start_row=21 + i, start_column=1, end_row=21 + i, end_column=5)
    c.alignment = LW
    w4.row_dimensions[21 + i].height = 32

# ===================================================================== Сравнение методик
w5 = wb.create_sheet("Сравнение методик", 4)
widths(w5, [20, 11, 11, 11, 12, 13, 13, 12, 12, 14])
title(w5, "4. Зарубежная методика (схема подбора, семинар 1, стр. 17) и сравнение",
      "Цепочка от 5\": 6 1/8\" (уменьш. зазор) → 7 5/8\" → 9 1/2\" → 10 3/4\" → 12 1/4\" (уменьш. зазор)")
head(w5, 4, 1, ["Колонна", "ОК, дюйм", "ОК, мм", "долото, дюйм", "долото, мм", "ОК (РФ), мм",
                "долото (РФ), мм", "Δ факт., мм", "Δ норм., мм", "тип зазора"])
colmap = {"Кондуктор": "C", "Промежуточная": "D", "Эксплуатационная": "E"}
for i, (n, ci, cm, bi, bm, kind) in enumerate(FOREIGN):
    r = 5 + i
    c = colmap[n]
    w5.cell(r, 1, n); w5.cell(r, 2, ci); w5.cell(r, 3, cm); w5.cell(r, 4, bi); w5.cell(r, 5, bm)
    w5.cell(r, 6, f"=Конструкция!{c}10")
    w5.cell(r, 7, f"=Конструкция!{c}17")
    w5.cell(r, 8, f"=(E{r}-INDEX({CM},MATCH(C{r},{CD},0)))/2")
    w5.cell(r, 9, f"=Конструкция!{c}15")
    w5.cell(r, 10, kind)
fmt(w5, "A5:J7", nf="0.0")
fmt(w5, "C5:C7", RES, nf="0.0")
fmt(w5, "E5:E7", RES, nf="0.0")
w5["A9"] = "Совпадают ли диаметры колонн?"
w5["F9"] = '=IF(AND(C5=F5,C6=F6,C7=F7),"да — все три колонны одинаковые","нет")'
w5["A10"] = "Разница диаметров долот (заруб. − РФ), мм:"
w5["F10"] = '="экспл.: "&FIXED(E5-G5,1)&"; пром.: "&FIXED(E6-G6,1)&"; конд.: "&FIXED(E7-G7,1)'
for r in (9, 10):
    w5.merge_cells(f"A{r}:E{r}")
    w5.merge_cells(f"F{r}:J{r}")
fmt(w5, "A9:J10", align=LW)
w5["A12"] = ("Вывод: обе методики дают одинаковые колонны 127,0 / 193,7 / 273,1 мм. Зарубежная схема допускает "
             "меньшие зазоры, поэтому долота меньше: 155,6 (нет в ГОСТ 20692) вместо 161,0; 241,3 вместо 244,5; "
             "311,2 вместо 349,2 мм. Для 5\" и 10 3/4\" колонн схема даёт только вариант с уменьшенным зазором.")
w5.merge_cells("A12:J14")
w5["A12"].alignment = LW

for w in wb.worksheets:
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
