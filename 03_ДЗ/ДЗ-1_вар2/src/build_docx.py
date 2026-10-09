"""Пояснительная записка ДЗ-1 (Word): Markdown + LaTeX -> pandoc (формулы — редактор формул Word) -> оформление python-docx."""
import os
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from calc import (SUB, LAYERS, RU, CEMENT, SECTIONS, section_design, COND_DEPTH, INTER_DEPTH,
                  PROD_DEPTH, G, RHO_W, WALL_CHOICE, OTTM, SMALL_DELTA)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
FIG = os.path.join(ROOT, "рисунки")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "ДЗ1_вар2_пояснительная_записка.docx")

STUDENT = "Фамилия И.О."
GROUP = "РТ\u201123\u201112"


def f(x, nd=3):
    return f"{x:.{nd}f}".replace(".", ",")


def m(x, nd=1):
    return f"{x:.{nd}f}".replace(".", "{,}")  # для LaTeX


PB = "\\newpage"  # заменяется на разрыв страницы

L = []
add = L.append

# ------------------------------------------------------------------ титул
add('::: {custom-style="Титул"}')
add("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ\n")
add("РОССИЙСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ НЕФТИ И ГАЗА\n(НАЦИОНАЛЬНЫЙ ИССЛЕДОВАТЕЛЬСКИЙ УНИВЕРСИТЕТ)\nИМЕНИ И.М. ГУБКИНА\n")
add("Кафедра бурения нефтяных и газовых скважин\n")
add(":::")
add('::: {custom-style="ТитулОтступ"}\n \n:::')
add('::: {custom-style="ТитулЗаголовок"}')
add("ДОМАШНЕЕ ЗАДАНИЕ № 1\n")
add(":::")
add('::: {custom-style="Титул"}')
add("по дисциплине «Технология бурения скважин»\n")
add("«Проектирование конструкции скважины»\n")
add("Вариант 2\n")
add(":::")
add('::: {custom-style="ТитулОтступ"}\n \n:::')
add('::: {custom-style="ТитулПодписи"}')
add(f"Выполнил: студент гр. {GROUP}\n")
add(f"[[{STUDENT}]]\n")
add(" \n")
add("Проверил: доцент, к.т.н.\n")
add("Балицкий В.П.\n")
add(":::")
add('::: {custom-style="ТитулОтступ"}\n \n:::')
add('::: {custom-style="Титул"}')
add("Москва, 2026")
add(":::")
add(PB)

# ------------------------------------------------------------------ 1
add("# Исходные данные")
add("Требуется спроектировать конструкцию газовой скважины глубиной 2400 м с эксплуатационной колонной "
    "диаметром 127,0 мм. До глубины 500 м разрез сложен неустойчивыми породами, склонными к поглощениям. "
    "Литология и давления приведены в таблице 1 (давления отнесены к подошве интервала).")
add('::: {custom-style="ПодписьТаблицы"}\nТаблица 1 – Геологический разрез\n:::')
add("| Интервал залегания, м | Порода | $P_F$, МПа | $P_{FR}$, МПа |")
add("|:---:|:---:|:---:|:---:|")
for t, b, rock, pf, pfr in LAYERS:
    add(f"| {t}–{b} | {rock} | {f(pf,1)} | {f(pfr,1)} |")
add("")
add("Расчёты выполнены в Excel (прилагаемый файл); ниже приведены расчётные зависимости, результаты "
    "и обоснование принятых решений.")

# ------------------------------------------------------------------ 2
add("# Индексы давлений и плотность бурового раствора")
add("Индексы пластового давления и давления поглощения:")
add(r"$$K_F = \frac{P_F}{\rho_w g z}, \qquad K_{FR} = \frac{P_{FR}}{\rho_w g z}, \qquad (1)$$")
add(r"где $P_F$, $P_{FR}$ — пластовое давление и давление поглощения на глубине $z$, Па; "
    r"$\rho_w = 1000$ кг/м³; $g = 9{,}81$ м/с².")
add("Относительная плотность бурового раствора выбирается с превышением гидростатического давления "
    "над пластовым не менее 10 % до глубины 1200 м и 5 % глубже [4, п. 2.7.3.3]:")
add(r"$$\rho_R = K_S K_F, \qquad (2)$$")
add(r"где $K_S = 1{,}1$ при $z \le 1200$ м и $K_S = 1{,}05$ при $z > 1200$ м. Плотность раствора "
    r"$\rho_m = \rho_R \rho_w$.")
s2 = SUB[2]
add("Известняк (1100–2100 м) пересекает глубину 1200 м, поэтому для него $\\rho_R$ определена отдельно "
    "на участках 1100–1200 и 1200–2100 м. Расчёт для участка 1100–1200 м ($z = 2100$ м — подошва пласта):")
add(rf"$$K_F = \frac{{27{{,}}3 \cdot 10^6}}{{1000 \cdot 9{{,}}81 \cdot 2100}} = {m(s2['kf'],3)}; \quad "
    rf"K_{{FR}} = \frac{{33{{,}}6 \cdot 10^6}}{{1000 \cdot 9{{,}}81 \cdot 2100}} = {m(s2['kfr'],3)}; \quad "
    rf"\rho_R = 1{{,}}1 \cdot {m(s2['kf'],3)} = {m(s2['rr'],3)}.$$")
add('::: {custom-style="ПодписьТаблицы"}\nТаблица 2 – Индексы давлений и плотность бурового раствора\n:::')
add("| Интервал, м | Порода | $K_F$ | $K_{FR}$ | $K_S$ | $\\rho_R$ | $\\rho_m$, кг/м³ |")
add("|:---:|:---:|:---:|:---:|:---:|:---:|:---:|")
for s in SUB:
    add(f"| {s['top']}–{s['bot']} | {s['rock']} | {f(s['kf'])} | {f(s['kfr'])} | {f(s['ks'],2)} | {f(s['rr'])} | {f(s['rr']*1000,0)} |")
add("")
add("В интервале 900–1100 м пластовое давление ниже гидростатического ($K_F = 0{,}816$), "
    "поэтому расчётная $\\rho_R < 1$; этот интервал бурится раствором вышележащих пород.")

# ------------------------------------------------------------------ 3
d1 = section_design(0, COND_DEPTH)
d2 = section_design(COND_DEPTH, INTER_DEPTH)
d3 = section_design(INTER_DEPTH, PROD_DEPTH)
add("# Совмещённый график давлений")
add("Совмещённый график давлений (рисунок 1) построен по данным таблицы 2. Закрашенная область между "
    "$K_F$ и $K_{FR}$ — диапазон допустимых плотностей раствора. Бурение интервалов одним открытым стволом "
    "возможно, если одна плотность раствора удовлетворяет условию")
add(r"$$K_F < \rho_R < K_{FR} \qquad (3)$$")
add("во всех вскрытых интервалах; в противном случае вышележащий интервал перекрывается обсадной колонной.")
add("![](рисунки/рис1_график_совмещенных_давлений.png){width=14cm}")
add('::: {custom-style="ПодписьРисунка"}\nРисунок 1 – Совмещённый график давлений\n:::')
add("Для интервала 1100–2400 м требуемая плотность определяется песчаником: "
    f"$\\rho_R = {m(d3['rr'],3)}$. Она больше индекса давления поглощения вышележащих пород "
    f"($K_{{FR}} = {m(SUB[0]['kfr'],3)}$ и ${m(SUB[1]['kfr'],3)}$), а плотность, допустимая для интервала 0–1100 м "
    f"($\\rho_R = {m(d2['rr'],3)}$), меньше $K_F = {m(SUB[2]['kf'],3)}$ известняка. Следовательно, условия бурения "
    "выше и ниже 1100 м несовместимы (таблица 3).")
add('::: {custom-style="ПодписьТаблицы"}\nТаблица 3 – Проверка совместимости условий бурения\n:::')
add("| Интервал открытого ствола, м | $\\rho_R$ принятая | $K_{F\\,max}$ | $K_{FR\\,min}$ | Вывод |")
add("|:----------:|:-----:|:-----:|:-----:|:------:|")
add(f"| 0–2400 (без промежуточной колонны) | {f(d3['rr'])} | {f(d3['kf_max'])} | {f(SUB[0]['kfr'])} | несовместимы |")
add(f"| 0–1100 | {f(d2['rr'])} | {f(d2['kf_max'])} | {f(d2['kfr_min'])} | совместимы |")
add(f"| 1100–2400 | {f(d3['rr'])} | {f(d3['kf_max'])} | {f(d3['kfr_min'])} | совместимы |")
add("")

# ------------------------------------------------------------------ 4
add("# Конструкция скважины")
add("## Число и глубины спуска обсадных колонн")
add("Принята конструкция из трёх обсадных колонн:")
add(f"- **кондуктор — {COND_DEPTH} м**: перекрывает неустойчивые породы, склонные к поглощениям (0–500 м); "
    "башмак устанавливается на 10 м ниже их подошвы;")
add(f"- **промежуточная колонна — {INTER_DEPTH} м**: изолирует интервал 0–1100 м, несовместимый с нижележащими "
    "известняком и песчаником с аномально высоким пластовым давлением. Башмак — в подошве алевролитов; "
    f"углубление в известняк раствором $\\rho_R = {m(d2['rr'],3)} < K_F = {m(SUB[2]['kf'],3)}$ вызвало бы газопроявление;")
add(f"- **эксплуатационная колонна — {PROD_DEPTH} м** (до проектной глубины), диаметр 127,0 мм задан.")
add("")
add("## Интервалы цементирования")
add("По Правилам безопасности [4, п. 2.7.4.10–2.7.4.11] кондуктор цементируется до устья, а в газовых "
    "скважинах тампонажный раствор поднимается не менее чем на 500 м над башмаком предыдущей колонны "
    "и кровлей продуктивного горизонта.")
add('::: {custom-style="ПодписьТаблицы"}\nТаблица 4 – Интервалы цементирования\n:::')
add("| Колонна | Башмак, м | Расчётная кровля цемента, м | Принятый интервал, м |")
add("|:---:|:---:|:---:|:---:|")
add(f"| Кондуктор | {COND_DEPTH} | до устья | 0–{COND_DEPTH} |")
add(f"| Промежуточная | {INTER_DEPTH} | {COND_DEPTH} − 500 = 10 | 0–{INTER_DEPTH} |")
add(f"| Эксплуатационная | {PROD_DEPTH} | {INTER_DEPTH} − 500 = 600 | 600–{PROD_DEPTH} |")
add("")
add("Для промежуточной колонны расчётная кровля цемента (10 м) практически совпадает с устьем, поэтому "
    "колонна цементируется до устья. Цементное кольцо эксплуатационной колонны перекрывает продуктивный "
    "песчаник (2100–2400 м) с запасом 1500 м над его кровлей и известняк с АВПД; перекрытие с промежуточной "
    "колонной — 500 м.")
add("## Диаметры обсадных колонн и долот")
add("Расчёт ведётся от эксплуатационной колонны вверх. Диаметр долота под колонну:")
add(r"$$D_B = d_m + 2\Delta, \qquad (4)$$")
add(r"где $d_m$ — наружный диаметр муфты (трубы ОТТМ по ГОСТ 632-80 [2]); $\Delta$ — минимальный "
    r"радиальный зазор [1]: 8 мм для колонн 114–127 мм, 13 мм — 168–245 мм, 18 мм — 273–299 мм. "
    r"Принимается ближайшее долото по ГОСТ 20692-2003 [3], не меньшее расчётного. Предыдущая колонна "
    r"должна пропускать это долото:")
add(r"$$d_i \ge D_B + 2\delta, \qquad (5)$$")
add(rf"где $\delta = {m(SMALL_DELTA,0)}$ мм — зазор между долотом и внутренней стенкой колонны (принят из диапазона 5–10 мм).")
rp, ri, rc = RU
add(rf"Эксплуатационная колонна 127,0 мм ($d_m = 141{{,}}3$ мм, $\Delta = 8$ мм): "
    rf"$D_B = 141{{,}}3 + 2 \cdot 8 = {m(rp['db_calc'])}$ мм → долото **{f(rp['db'],1)} мм**.")
add(rf"Промежуточная колонна: $d_i \ge {m(rp['db'])} + 2 \cdot 5 = {m(ri['di_min_this'])}$ мм. Колонны 168,3 и 177,8 мм "
    "не подходят (наибольший внутренний диаметр 153,7 и 164,0 мм). Принята колонна 193,7 мм, допустимая "
    "толщина стенки 7,6…10,9 мм; предварительно принята 10,9 мм (d = 171,9 мм), окончательно — по расчёту "
    rf"на прочность. $D_B = 215{{,}}9 + 2 \cdot 13 = {m(ri['db_calc'])}$ мм → долото **{f(ri['db'],1)} мм**.")
add(rf"Кондуктор: $d_i \ge {m(ri['db'])} + 2 \cdot 5 = {m(rc['di_min_this'])}$ мм. Колонна 244,5 мм не подходит "
    "(d ≤ 228,7 мм); у колонны 273,1 мм условию удовлетворяет только толщина стенки 8,9 мм (d = 255,3 мм). "
    rf"$D_B = 298{{,}}5 + 2 \cdot 18 = {m(rc['db_calc'])}$ мм → долото **{f(rc['db'],1)} мм**.")
add('::: {custom-style="ПодписьТаблицы"}\nТаблица 5 – Диаметры обсадных колонн и долот\n:::')
add("| Колонна | $D$, мм | $s$, мм | $d_m$, мм | $\\Delta$, мм | $D_B$ расч., мм | $D_B$ ГОСТ, мм |")
add("|:----------------:|:-----:|:-----:|:-----:|:-----:|:-------:|:-------:|")
add(f"| Кондуктор | {f(rc['D'],1)} | {f(WALL_CHOICE['Кондуктор'],1)} | {f(rc['dm'],1)} | {rc['delta']} | {f(rc['db_calc'],1)} | **{f(rc['db'],1)}** |")
add(f"| Промежуточная | {f(ri['D'],1)} | {f(WALL_CHOICE['Промежуточная'],1)} | {f(ri['dm'],1)} | {ri['delta']} | {f(ri['db_calc'],1)} | **{f(ri['db'],1)}** |")
add(f"| Эксплуатационная | {f(rp['D'],1)} | — | {f(rp['dm'],1)} | {rp['delta']} | {f(rp['db_calc'],1)} | **{f(rp['db'],1)}** |")
add("")
add("## Схема конструкции скважины")
add("Схема конструкции выполнена в обозначениях [1] (рисунок 2): D — диаметр колонны, L — глубина спуска, "
    "d — диаметр долота, H — кровля цементного камня.")
add("![](рисунки/рис2_схема_конструкции.png){width=11.5cm}")
add('::: {custom-style="ПодписьРисунка"}\nРисунок 2 – Схема конструкции скважины\n:::')

# ------------------------------------------------------------------ 5
add("# Сравнение с зарубежной методикой")
add("По зарубежной методике диаметры подбираются по схеме соответствия колонн и долот [1] (рисунок 3); "
    "чёрные стрелки — стандартный зазор, серые — уменьшенный.")
add("![](рисунки/рис3_зарубежная_схема_выбор.png){width=16.5cm}")
add('::: {custom-style="ПодписьРисунка"}\nРисунок 3 – Подбор диаметров по зарубежной схеме (выбранная цепочка выделена)\n:::')
add("Для колонны 5\" (127,0 мм) схема даёт только долото с уменьшенным зазором 6 1/8\" (155,6 мм); далее: колонна "
    "7 5/8\" (193,7 мм) → долото 9 1/2\" (241,3 мм) → колонна 10 3/4\" (273,1 мм) → долото 12 1/4\" (311,2 мм), "
    "последний переход — также только с уменьшенным зазором. Сравнение — в таблице 6.")
add('::: {custom-style="ПодписьТаблицы"}\nТаблица 6 – Результаты двух методик\n:::')
add("| Колонна | Колонна (РФ / заруб.), мм | Долото РФ, мм | Долото заруб., мм | $\\Delta$ заруб., мм | $\\Delta$ норм., мм |")
add("|:---------------:|:-------------:|:------:|:----------:|:------:|:------:|")
add("| Эксплуатационная | 127,0 / 127,0 (5\") | 161,0 | 155,6 (6 1/8\") | 7,2 | 8 |")
add("| Промежуточная | 193,7 / 193,7 (7 5/8\") | 244,5 | 241,3 (9 1/2\") | 12,7 | 13 |")
add("| Кондуктор | 273,1 / 273,1 (10 3/4\") | 349,2 | 311,2 (12 1/4\") | 6,4 | 18 |")
add("")
add("Диаметры всех трёх обсадных колонн по обеим методикам совпадают, т.е. зарубежная методика даёт "
    "аналогичный результат. Различаются диаметры долот: по зарубежной схеме они меньше (на 5,4; 3,2 и 38,0 мм), "
    "так как она допускает меньшие радиальные зазоры между муфтой и стенкой скважины — 7,2; 12,7 и 6,4 мм "
    "против нормативных 8, 13 и 18 мм. Долота 155,6 мм нет в ряду ГОСТ 20692-2003 (ближайшие 151,0 и 161,0 мм). "
    "Зарубежный вариант уменьшает объём выбуриваемой породы, особенно под кондуктор; отечественный за счёт "
    "больших зазоров облегчает спуск колонн и повышает качество цементирования, что важно для газовой "
    "скважины с АВПД.")

# ------------------------------------------------------------------ заключение
add("# ЗАКЛЮЧЕНИЕ {.unnumbered}")
add("1. Рассчитаны индексы давлений; плотность бурового раствора: $\\rho_R = "
    f"{m(d2['rr'],2)}$ в интервале 0–1100 м и $\\rho_R = {m(d3['rr'],2)}$ в интервале 1100–2400 м.")
add("2. По совмещённому графику давлений установлена несовместимость условий бурения выше и ниже 1100 м.")
add(f"3. Принята конструкция: кондуктор 273,1 мм — {COND_DEPTH} м (долото 349,2 мм, цемент до устья); "
    f"промежуточная колонна 193,7 мм — {INTER_DEPTH} м (долото 244,5 мм, цемент до устья); эксплуатационная "
    f"колонна 127,0 мм — {PROD_DEPTH} м (долото 161,0 мм, цемент 600–2400 м).")
add("4. Зарубежная методика даёт те же диаметры обсадных колонн при меньших диаметрах долот.")

# ------------------------------------------------------------------ литература
add("# СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ {.unnumbered}")
add('::: {custom-style="Литература"}')
for i, t in enumerate([
    "Проектирование конструкции скважины: материалы практического занятия № 1 по дисциплине «Технология бурения скважин». — М.: РГУ нефти и газа (НИУ) имени И.М. Губкина, кафедра бурения нефтяных и газовых скважин.",
    "ГОСТ 632-80. Трубы обсадные и муфты к ним. Технические условия.",
    "ГОСТ 20692-2003. Долота шарошечные. Технические условия.",
    "Правила безопасности в нефтяной и газовой промышленности (ПБ 08-624-03).",
], 1):
    add(f"{i}. {t}")
add(":::")


parts = []
for line in L:
    if parts and line.startswith("|") and parts[-1].startswith("|"):
        parts[-1] += "\n" + line
    else:
        parts.append(line)
md = "\n\n".join(parts)
md = md.replace(PB, '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```')
md = md.replace("[[TOC]]", '```{=openxml}\n<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
                '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Для обновления содержания нажмите F9 (правой кнопкой → Обновить поле).</w:t></w:r>'
                '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>\n```')

tmp = tempfile.mkdtemp()
md_path = os.path.join(tmp, "note.md")
open(md_path, "w", encoding="utf-8").write(md)
raw = os.path.join(tmp, "raw.docx")
subprocess.run(["pandoc", md_path, "-f", "markdown+pipe_tables+tex_math_dollars+raw_attribute+fenced_divs",
                "-o", raw, "--resource-path", ROOT, "--number-sections"], check=True)

# ================================================================== оформление
doc = Document(raw)
TNR = "Times New Roman"


def set_font(style, size=14, bold=None, italic=None):
    style.font.name = TNR
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), TNR)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]


def para_fmt(style, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25, before=0, after=0, spacing=1.5, keep_next=False):
    pf = style.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(indent) if indent else Cm(0)
    pf.left_indent = Cm(0)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    pf.keep_with_next = keep_next


class _Styles:
    def __init__(self, st):
        self._st = st

    def __iter__(self):
        return iter(self._st)

    def __getitem__(self, name):
        for x in self._st:
            if x.name == name:
                return x
        raise KeyError(name)


styles = _Styles(doc.styles)
for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
    if name in [s.name for s in styles]:
        set_font(styles[name])
        para_fmt(styles[name])
for name, lvl in (("Heading 1", 1), ("Heading 2", 2)):
    st = styles[name]
    set_font(st, 14, bold=True, italic=False)
    para_fmt(st, WD_ALIGN_PARAGRAPH.LEFT, indent=1.25, before=12, after=12, keep_next=True)
custom = {
    "Титул": dict(align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=False),
    "ТитулЗаголовок": dict(align=WD_ALIGN_PARAGRAPH.CENTER, size=20, bold=True),
    "ТитулПодписи": dict(align=WD_ALIGN_PARAGRAPH.LEFT, size=14, bold=False, left=8.0),
    "ТитулОтступ": dict(align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=False, before=60),
    "ПодписьТаблицы": dict(align=WD_ALIGN_PARAGRAPH.LEFT, size=14, bold=False, before=6, keep=True),
    "ПодписьРисунка": dict(align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=False, after=6),
    "Литература": dict(align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, bold=False),
}
for name, o in custom.items():
    if name not in [s.name for s in styles]:
        continue
    st = styles[name]
    set_font(st, o["size"], bold=o["bold"])
    para_fmt(st, o["align"], indent=0, before=o.get("before", 0), after=o.get("after", 0),
             keep_next=o.get("keep", False))
    if o.get("left"):
        st.paragraph_format.left_indent = Cm(o["left"])
if "Литература" in [s.name for s in styles]:
    styles["Литература"].paragraph_format.first_line_indent = Cm(1.25)

# поля и колонтитул с номером страницы (на титуле номера нет)
for sec in doc.sections:
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Cm(0)
    run = fp.add_run()
    for t, txt in (("begin", None), (None, " PAGE "), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt; run._r.append(el)
    run.font.name = TNR
    run.font.size = Pt(12)

# абзацы с формулами и рисунками — по центру, без отступа; списки — без лишних отступов
for p in doc.paragraphs:
    xml = p._p.xml
    if "<m:oMathPara" in xml or "<w:drawing" in xml:
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "<w:drawing" in xml:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(6)
    # выделить заглушку ФИО
    for r in p.runs:
        if "[[" in r.text:
            r.text = r.text.replace("[[", "").replace("]]", "")
            hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow")
            r._r.get_or_add_rPr().append(hl)

# таблицы: сетка, 12 пт, одинарный интервал, по центру
def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(b)


for tbl in doc.tables:
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(tbl)
    tblW = tbl._tbl.tblPr.find(qn("w:tblW"))
    if tblW is not None:
        tblW.set(qn("w:type"), "pct"); tblW.set(qn("w:w"), "5000")
    for ri, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                if p.paragraph_format.alignment is None:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(12)
                    r.font.name = TNR
                    if ri == 0:
                        r.font.bold = True

# строки таблиц не разрываются, таблица держится целиком
for tbl in doc.tables:
    n = len(tbl.rows)
    for ri, row in enumerate(tbl.rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if ri < n - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True

# отступ после таблиц
for tbl in doc.tables:
    nxt = tbl._tbl.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        pPr = nxt.get_or_add_pPr()
        sp = pPr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing"); pPr.append(sp)
        sp.set(qn("w:before"), "160")

doc.core_properties.author = STUDENT
doc.core_properties.title = "ДЗ-1. Проектирование конструкции скважины. Вариант 2"
doc.save(OUT)
print("saved", OUT)


# ---- десятичные числа в формулах: «9», «,», «81» -> одно число «9,81» без математического пробела после запятой
def _fix_decimal_commas(path):
    import re
    import shutil
    import zipfile
    tmp_path = path + ".tmp"
    pat = re.compile(r'<m:r>(?:<m:rPr>(?:<m:sty m:val="p"/>)?</m:rPr>)?<m:t>(\d+)</m:t></m:r>'
                     r'<m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>,</m:t></m:r>'
                     r'<m:r>(?:<m:rPr>(?:<m:sty m:val="p"/>)?</m:rPr>)?<m:t>(\d+)</m:t></m:r>')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                xml, n = pat.subn(r'<m:r><m:rPr><m:nor/></m:rPr><m:t>\1,\2</m:t></m:r>', xml)
                print("decimal commas fixed:", n)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp_path, path)


_fix_decimal_commas(OUT)
