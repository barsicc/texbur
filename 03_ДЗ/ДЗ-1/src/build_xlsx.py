"""Excel-файл с расчётами ДЗ-1 (все величины — формулами)."""
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from calc import LAYERS, OTTM, BITS, delta_min, FOREIGN, WALL_CHOICE

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "..", "ДЗ1_вар7_расчеты.xlsx")

wb = Workbook()
thin = Side(style="thin", color="000000")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="D9E1F2")
INPUT = PatternFill("solid", fgColor="FFF2CC")
RESULT = PatternFill("solid", fgColor="E2EFDA")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def box(ws, rng, fill=None, bold=False, align=CENTER, fmt=None):
    for row in ws[rng]:
        for c in row:
            c.border = BOX
            c.alignment = align
            if fill:
                c.fill = fill
            if bold:
                c.font = BOLD
            if fmt:
                c.number_format = fmt


def amin(ws, cell, vals, crit, x):
    """Минимум vals при crit >= x (формула массива, работает в любой версии Excel)."""
    ws[cell] = ArrayFormula(cell, f"=MIN(IF({crit}>={x},{vals}))")


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(bold=True, size=13)


# =====================================================================
# Лист 1. Исходные данные
# =====================================================================
ws = wb.active
ws.title = "1_Исходные данные"
widths(ws, [3, 14, 14, 18, 14, 14, 4, 40, 12, 12])
title(ws, "B1", "ДЗ-1. Проектирование конструкции скважины. Вариант 7")
ws["B3"] = "Интервал залегания, м"; ws.merge_cells("B3:C3")
ws["D3"] = "Порода"; ws.merge_cells("D3:D4")
ws["E3"] = "Давление, МПа"; ws.merge_cells("E3:F3")
ws["B4"], ws["C4"], ws["E4"], ws["F4"] = "кровля", "подошва", "пластовое Pпл", "поглощения Pпогл"
box(ws, "B3:F4", HEAD, True)
for i, (t, b, rock, pf, pfr) in enumerate(LAYERS):
    r = 5 + i
    ws.cell(r, 2, t); ws.cell(r, 3, b); ws.cell(r, 4, rock); ws.cell(r, 5, pf); ws.cell(r, 6, pfr)
box(ws, "B5:F8", INPUT)
ws["B9"] = "Давления отнесены к подошве интервала залегания."
ws["B9"].font = Font(italic=True, size=9)

params = [
    ("Назначение скважины", "газовая", ""),
    ("Глубина скважины H, м", 2100, ""),
    ("Диаметр эксплуатационной колонны, мм", 114.3, ""),
    ("Неустойчивые породы до глубины, м", 250, ""),
    ("Ускорение свободного падения g, м/с²", 9.81, ""),
    ("Плотность воды ρw, кг/м³", 1000, ""),
    ("Ks при z < 1200 м", 1.1, "ПБ п. 2.7.3.3"),
    ("Ks при z > 1200 м", 1.05, "ПБ п. 2.7.3.3"),
    ("Граница для Ks, м", 1200, ""),
    ("Подъём цемента над башмаком пред. ОК (газ), м", 500, "ПБ п. 2.7.4.11"),
    ("Зазор δ (долото – внутр. стенка ОК), мм", 5, "δ = 5–10 мм"),
    ("Допустимая репрессия (скв. глубже 1200 м), МПа", 2.5, "ПБ п. 2.7.3.3: 2,5–3,0"),
]
ws["H3"] = "Параметр"; ws["I3"] = "Значение"; ws["J3"] = "Источник"
box(ws, "H3:J3", HEAD, True)
P = {}
for i, (name, val, src) in enumerate(params):
    r = 4 + i
    ws.cell(r, 8, name); ws.cell(r, 9, val); ws.cell(r, 10, src)
    P[name] = f"'1_Исходные данные'!$I${r}"
box(ws, f"H4:J{3 + len(params)}", None, align=LEFT)
box(ws, f"I4:I{3 + len(params)}", INPUT)
G = P["Ускорение свободного падения g, м/с²"]
RW = P["Плотность воды ρw, кг/м³"]
KS1 = P["Ks при z < 1200 м"]
KS2 = P["Ks при z > 1200 м"]
ZKS = P["Граница для Ks, м"]
CEM = P["Подъём цемента над башмаком пред. ОК (газ), м"]
DELTA_S = P["Зазор δ (долото – внутр. стенка ОК), мм"]
REP = P["Допустимая репрессия (скв. глубже 1200 м), МПа"]
HW = P["Глубина скважины H, м"]
DPROD = P["Диаметр эксплуатационной колонны, мм"]
UNST = P["Неустойчивые породы до глубины, м"]

# =====================================================================
# Лист 2. Индексы давлений и плотность раствора
# =====================================================================
ws2 = wb.create_sheet("2_Индексы давлений")
widths(ws2, [3, 10, 10, 26, 13, 11, 11, 11, 10, 12, 13, 18])
title(ws2, "B1", "Расчёт индексов давлений и относительной плотности бурового раствора")
ws2["B2"] = "KF = Pпл/(ρw·g·z);  KFR = Pпогл/(ρw·g·z);  ρR = Ks·KF;  условие: KF < ρR < KFR"
ws2["B2"].font = Font(italic=True)
hdr = ["кровля, м", "подошва, м", "порода", "Pпл, МПа", "Pпогл, МПа", "KF", "KFR", "Ks", "ρR", "ρб.р., кг/м³", "KF<ρR<KFR"]
for j, h in enumerate(hdr):
    ws2.cell(4, 2 + j, h)
box(ws2, "B4:L4", HEAD, True)
# подынтервалы: (кровля, подошва, индекс слоя в листе 1)
subs = [(0, 800, 0), (800, 1000, 1), (1000, 1200, 2), (1200, 1900, 2), (1900, 2100, 3)]
for i, (t, b, li) in enumerate(subs):
    r = 5 + i
    src = 5 + li
    ws2.cell(r, 2, t)
    ws2.cell(r, 3, b)
    ws2.cell(r, 4, f"='1_Исходные данные'!D{src}")
    ws2.cell(r, 5, f"='1_Исходные данные'!E{src}")
    ws2.cell(r, 6, f"='1_Исходные данные'!F{src}")
    zc = f"'1_Исходные данные'!C{src}"   # глубина отнесения давления — подошва слоя
    ws2.cell(r, 7, f"=E{r}*10^6/({RW}*{G}*{zc})")
    ws2.cell(r, 8, f"=F{r}*10^6/({RW}*{G}*{zc})")
    ws2.cell(r, 9, f"=IF(C{r}<={ZKS},{KS1},{KS2})")
    ws2.cell(r, 10, f"=I{r}*G{r}")
    ws2.cell(r, 11, f"=J{r}*{RW}")
    ws2.cell(r, 12, f'=IF(AND(G{r}<J{r},J{r}<H{r}),"выполняется","НЕ выполняется")')
box(ws2, "B5:L9")
box(ws2, "G5:J9", RESULT, fmt="0.000")
box(ws2, "K5:K9", RESULT, fmt="0")
ws2["B10"] = ("Песчаник (1000–1900 м) разбит границей 1200 м, т.к. Ks меняется с 1,1 на 1,05. "
              "KF и KFR внутри пласта постоянны (давления отнесены к подошве пласта, 1900 м).")
ws2["B10"].font = Font(italic=True, size=9)

# Интервалы бурения (секции открытого ствола)
ws2["B12"] = "Совместимость условий бурения по интервалам открытого ствола"
ws2["B12"].font = BOLD
hdr = ["кровля, м", "подошва, м", "интервал бурения", "ρR принятая = max ρR", "KF max", "KFR min",
       "ρб.р., кг/м³", "", "", "", "совместимость"]
for j, h in enumerate(hdr):
    ws2.cell(13, 2 + j, h)
box(ws2, "B13:L13", HEAD, True)
sections = [
    (0, "='3_Колонны'!D5", "под кондуктор", "J5", "G5", "H5"),
    ("='3_Колонны'!D5", "='3_Колонны'!D6", "под промежуточную ОК", "MAX(J5:J6)", "MAX(G5:G6)", "MIN(H5:H6)"),
    ("='3_Колонны'!D6", "='3_Колонны'!D7", "под эксплуатационную ОК", "MAX(J7:J9)", "MAX(G7:G9)", "MIN(H7:H9)"),
]
for i, (a, b, name, rr, kf, kfr) in enumerate(sections):
    r = 14 + i
    ws2.cell(r, 2, a); ws2.cell(r, 3, b); ws2.cell(r, 4, name)
    ws2.cell(r, 5, f"={rr}"); ws2.cell(r, 6, f"={kf}"); ws2.cell(r, 7, f"={kfr}")
    ws2.cell(r, 8, f"=E{r}*{RW}")
    ws2.cell(r, 12, f'=IF(AND(F{r}<E{r},E{r}<G{r}),"совместимы","НЕсовместимы")')
box(ws2, "B14:L16")
box(ws2, "E14:G16", RESULT, fmt="0.000")
box(ws2, "H14:H16", RESULT, fmt="0")

ws2["B18"] = "Проверка необходимости промежуточной колонны (можно ли бурить 0–2100 м одним стволом)"
ws2["B18"].font = BOLD
ws2["B19"] = "ρR, требуемая для 1000–2100 м"; ws2["H19"] = "=E16"
ws2["B20"] = "KFR min в интервале 0–1000 м"; ws2["H20"] = "=MIN(H5:H6)"
ws2["B21"] = "ρR, принятая для 0–1000 м"; ws2["H21"] = "=E15"
ws2["B22"] = "KF в песчанике (1000–1900 м)"; ws2["H22"] = "=G7"
ws2["B23"] = "Вывод"
ws2["H23"] = ('=IF(OR(H19>=H20,H21<=H22),"несовместимы: интервал 0–1000 м перекрыть промежуточной колонной '
              'до вскрытия песчаника","совместимы")')
for r in range(19, 24):
    ws2.merge_cells(f"B{r}:G{r}")
box(ws2, "B19:H23", align=LEFT)
box(ws2, "H19:H22", RESULT, fmt="0.000")
ws2.merge_cells("H23:L23")
ws2.row_dimensions[23].height = 32

# Проверка репрессий и поглощений
ws2["B25"] = "Проверка репрессии и отсутствия поглощения (на подошве каждого подынтервала, принятая ρR)"
ws2["B25"].font = BOLD
hdr = ["z, м", "ρR принятая", "Pпл(z), МПа", "Pгс(z), МПа", "Pпогл(z), МПа", "ΔPрепр, МПа",
       "репрессия, %", "треб. ≥ %", "ΔP ≤ доп.", "Pгс < Pпогл", "итог"]
for j, h in enumerate(hdr):
    ws2.cell(26, 2 + j, h)
box(ws2, "B26:L26", HEAD, True)
sec_of = {0: 14, 1: 15, 2: 16, 3: 16, 4: 16}
for i in range(5):
    r = 27 + i
    sr = 5 + i
    ws2.cell(r, 2, f"=C{sr}")
    ws2.cell(r, 3, f"=E{sec_of[i]}")
    ws2.cell(r, 4, f"=G{sr}*{RW}*{G}*B{r}/10^6")
    ws2.cell(r, 5, f"=C{r}*{RW}*{G}*B{r}/10^6")
    ws2.cell(r, 6, f"=H{sr}*{RW}*{G}*B{r}/10^6")
    ws2.cell(r, 7, f"=E{r}-D{r}")
    ws2.cell(r, 8, f"=(E{r}/D{r}-1)*100")
    ws2.cell(r, 9, f"=(I{sr}-1)*100")
    ws2.cell(r, 10, f'=IF(G{r}<={REP},"да","нет")')
    ws2.cell(r, 11, f'=IF(E{r}<F{r},"да","нет")')
    ws2.cell(r, 12, f'=IF(AND(H{r}>=I{r}-0.01,J{r}="да",K{r}="да"),"ОК","проверить")')
box(ws2, "B27:L31")
box(ws2, "C27:C31", fmt="0.000")
box(ws2, "D27:G31", RESULT, fmt="0.00")
box(ws2, "H27:I31", fmt="0.0")
ws2["B32"] = ("Репрессия в 1000–1200 м = 10 % (ровно норма), ниже 1200 м — больше 5 %; "
              "максимальное превышение над пластовым — у подошвы песчаника (1900 м) — не превышает 2,5 МПа.")
ws2["B32"].font = Font(italic=True, size=9)

# Данные для графика (ступенчатые линии)
ws2["N3"] = "Данные для графика совмещённых давлений"; ws2["N3"].font = BOLD
for j, h in enumerate(["KF", "z", "KFR", "z", "ρR", "z", "ρR прин.", "z"]):
    ws2.cell(4, 14 + j, h)
box(ws2, "N4:U4", HEAD, True)
# KF/KFR — по слоям (строки листа 2: 5,6,7(=8),9)
layer_rows = [5, 6, 7, 9]
r = 5
for lr in layer_rows:
    top = f"B{lr}" if lr != 7 else "B7"
    bot = f"C{lr}" if lr != 7 else "C8"
    for zcell in (top, bot):
        ws2.cell(r, 14, f"=G{lr}"); ws2.cell(r, 15, f"={zcell}")
        ws2.cell(r, 16, f"=H{lr}"); ws2.cell(r, 17, f"={zcell}")
        r += 1
kf_last = r - 1
r = 5
for sr in range(5, 10):
    for zcell in (f"B{sr}", f"C{sr}"):
        ws2.cell(r, 18, f"=J{sr}"); ws2.cell(r, 19, f"={zcell}")
        r += 1
rr_last = r - 1
r = 5
for sr in (14, 15, 16):
    for zcell in (f"B{sr}", f"C{sr}"):
        ws2.cell(r, 20, f"=E{sr}"); ws2.cell(r, 21, f"={zcell}")
        r += 1
rp_last = r - 1
box(ws2, f"N5:U{rr_last}", fmt="0.000")

ch = ScatterChart()
ch.title = "Совмещённый график давлений"
ch.style = 13
ch.x_axis.title = "KF; KFR; ρR"
ch.y_axis.title = "Глубина z, м"
ch.y_axis.scaling.orientation = "maxMin"
ch.y_axis.scaling.min = 0
ch.y_axis.scaling.max = 2200
ch.y_axis.majorUnit = 200
ch.x_axis.scaling.min = 0.8
ch.x_axis.scaling.max = 1.7
ch.x_axis.majorUnit = 0.1
ch.x_axis.number_format = "0.0"
ch.y_axis.number_format = "0"
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.height = 16
ch.width = 18
for (xc, yc, last, name, color, dash, w) in [
    (14, 15, kf_last, "KF", "000000", None, 28575),
    (16, 17, kf_last, "KFR", "7F7F7F", None, 28575),
    (18, 19, rr_last, "ρR", "1F4E79", "dash", 19050),
    (20, 21, rp_last, "ρR принятая", "C00000", "dashDot", 12700),
]:
    xs = Reference(ws2, min_col=xc, min_row=5, max_row=last)
    ys = Reference(ws2, min_col=yc, min_row=5, max_row=last)
    s = Series(ys, xs, title=name)
    s.marker.symbol = "none"
    s.smooth = False
    s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=color, w=w, prstDash=dash))
    ch.series.append(s)
ws2.add_chart(ch, "N24")

# =====================================================================
# Лист 3. Колонны и цементирование
# =====================================================================
ws3 = wb.create_sheet("3_Колонны")
widths(ws3, [3, 22, 12, 12, 85])
title(ws3, "B1", "Число и глубины спуска обсадных колонн, интервалы цементирования")
for j, h in enumerate(["Колонна", "кровля, м", "башмак, м", "Обоснование глубины спуска"]):
    ws3.cell(4, 2 + j, h)
box(ws3, "B4:E4", HEAD, True)
rows3 = [
    ("Кондуктор", 0, f"={UNST}+10",
     "Перекрывает неустойчивые породы, склонные к поглощениям (0–250 м); башмак на 10 м ниже, в устойчивом алевролите."),
    ("Промежуточная", 0, "='1_Исходные данные'!C6",
     "Изолирует интервал 0–1000 м (KFR min = 1,22) перед вскрытием песчаника с АВПД (требуется ρR = 1,40). "
     "Башмак — в подошве глин (непроницаемые породы) на границе с песчаником."),
    ("Эксплуатационная", 0, f"={HW}",
     "До проектной глубины; Ø задан заданием (114,3 мм)."),
]
for i, (n, t, b, why) in enumerate(rows3):
    r = 5 + i
    ws3.cell(r, 2, n); ws3.cell(r, 3, t); ws3.cell(r, 4, b); ws3.cell(r, 5, why)
box(ws3, "B5:D7")
box(ws3, "D5:D7", RESULT)
box(ws3, "E5:E7", align=LEFT)
for r in range(5, 8):
    ws3.row_dimensions[r].height = 48

ws3["B9"] = "Интервалы цементирования"; ws3["B9"].font = BOLD
for j, h in enumerate(["Колонна", "кровля цемента, м", "подошва, м", "Обоснование"]):
    ws3.cell(10, 2 + j, h)
box(ws3, "B10:E10", HEAD, True)
cem = [
    ("Кондуктор", "=0", "=D5", "ПБ п. 2.7.4.10: кондуктор цементируется до устья."),
    ("Промежуточная", f"=MAX(0,D5-{CEM})", "=D6",
     "ПБ п. 2.7.4.11: подъём над башмаком предыдущей ОК не менее 500 м (газовая скв.): 260 − 500 < 0 → до устья."),
    ("Эксплуатационная", f"=MAX(0,D6-{CEM})", "=D7",
     "ПБ п. 2.7.4.11: не менее 500 м над башмаком промежуточной (1000 м) → кровля цемента 500 м; "
     "перекрывает продуктивный известняк, проницаемый песчаник с АВПД; над кровлей продуктивного пласта (1900 м) 1400 м > 500 м."),
]
for i, (n, a, b, why) in enumerate(cem):
    r = 11 + i
    ws3.cell(r, 2, n); ws3.cell(r, 3, a); ws3.cell(r, 4, b); ws3.cell(r, 5, why)
box(ws3, "B11:D13")
box(ws3, "C11:C13", RESULT)
box(ws3, "E11:E13", align=LEFT)
for r in range(11, 14):
    ws3.row_dimensions[r].height = 62
ws3["B14"] = "Высота цементного кольца эксплуатационной колонны, м"; ws3["D14"] = "=D13-C13"
ws3["B15"] = "Перекрытие с промежуточной колонной, м"; ws3["D15"] = "=D6-C13"
box(ws3, "B14:D15", align=LEFT)

# =====================================================================
# Лист 6/7. Справочники (создаём заранее для ссылок)
# =====================================================================
wsC = wb.create_sheet("ГОСТ 632-80 (ОТТМ)")
widths(wsC, [3, 12, 12, 10, 10, 14])
title(wsC, "B1", "Трубы обсадные ОТТМ и муфты (ГОСТ 632-80), мм")
for j, h in enumerate(["условный Ø", "D", "s", "d", "Dн муфты"]):
    wsC.cell(3, 2 + j, h)
box(wsC, "B3:F3", HEAD, True)
r = 4
for D in sorted(OTTM):
    nom, dm, walls = OTTM[D]
    for s, d in walls:
        wsC.cell(r, 2, nom); wsC.cell(r, 3, D); wsC.cell(r, 4, s); wsC.cell(r, 5, d); wsC.cell(r, 6, dm)
        r += 1
C_LAST = r - 1
box(wsC, f"B4:F{C_LAST}", fmt="0.0")
wsC["H3"] = "Минимальный радиальный зазор Δ (семинар 1)"; wsC["H3"].font = BOLD
wsC["H4"], wsC["I4"] = "условный Ø до", "Δ, мм"
box(wsC, "H4:I4", HEAD, True)
for i, (lim, dl) in enumerate([(127, 8), (146, 10), (245, 13), (299, 18), (426, 22)]):
    wsC.cell(5 + i, 8, lim); wsC.cell(5 + i, 9, dl)
box(wsC, "H5:I9")
wsC["H11"] = "Опечатки исходного docx исправлены по d = D − 2s."
wsC["H11"].font = Font(italic=True, size=9)

wsB = wb.create_sheet("ГОСТ 20692-2003 (долота)")
widths(wsB, [3, 16])
title(wsB, "B1", "Ряд диаметров шарошечных долот, мм")
wsB["B3"] = "D долота"; box(wsB, "B3:B3", HEAD, True)
for i, b in enumerate(BITS):
    wsB.cell(4 + i, 2, b)
B_LAST = 3 + len(BITS)
box(wsB, f"B4:B{B_LAST}", fmt="0.0")

CD = f"'ГОСТ 632-80 (ОТТМ)'!$C$4:$C${C_LAST}"
CS = f"'ГОСТ 632-80 (ОТТМ)'!$D$4:$D${C_LAST}"
Cd = f"'ГОСТ 632-80 (ОТТМ)'!$E$4:$E${C_LAST}"
CM = f"'ГОСТ 632-80 (ОТТМ)'!$F$4:$F${C_LAST}"
CN = f"'ГОСТ 632-80 (ОТТМ)'!$B$4:$B${C_LAST}"
BB = f"'ГОСТ 20692-2003 (долота)'!$B$4:$B${B_LAST}"
DLIM = "'ГОСТ 632-80 (ОТТМ)'!$H$5:$H$9"
DVAL = "'ГОСТ 632-80 (ОТТМ)'!$I$5:$I$9"

# =====================================================================
# Лист 4. Диаметры — отечественная методика
# =====================================================================
ws4 = wb.create_sheet("4_Диаметры (РФ)")
wb.move_sheet(ws4, offset=-2)
widths(ws4, [3, 46, 10, 18, 16, 14])
title(ws4, "B1", "Расчёт диаметров долот и обсадных колонн (отечественная методика, снизу вверх)")
ws4["B2"] = "DB = dм + 2Δ  →  долото по ГОСТ 20692 ≥ DB;   d пред. ОК ≥ DB + 2δ  →  ОК по ГОСТ 632-80"
ws4["B2"].font = Font(italic=True)
for j, h in enumerate(["Параметр", "обозн.", "Эксплуатационная", "Промежуточная", "Кондуктор"]):
    ws4.cell(4, 2 + j, h)
box(ws4, "B4:F4", HEAD, True)
labels = [
    ("Глубина спуска L, м", "L"),                                  # 5
    ("Требуемый внутр. диаметр d ≥ DB(след.) + 2δ, мм", "d min"),  # 6
    ("Наружный диаметр ОК D, мм", "D"),                             # 7
    ("Условный диаметр, мм", "Dусл"),                               # 8
    ("Толщина стенки s, мм (предварительно)", "s"),                 # 9
    ("Внутренний диаметр d, мм", "d"),                              # 10
    ("Наружный диаметр муфты ОТТМ dм, мм", "dм"),                   # 11
    ("Минимальный радиальный зазор Δ, мм", "Δ"),                    # 12
    ("Расчётный диаметр долота DB = dм + 2Δ, мм", "DB расч"),       # 13
    ("Диаметр долота по ГОСТ 20692 (≥ DB расч), мм", "DB"),         # 14
    ("Фактический радиальный зазор (DB − dм)/2, мм", "Δ факт"),     # 15
    ("Фактический зазор δ для долота след. интервала, мм", "δ факт"),  # 16
]
for i, (lab, sym) in enumerate(labels):
    ws4.cell(5 + i, 2, lab); ws4.cell(5 + i, 3, sym)
# L
ws4["D5"] = "='3_Колонны'!D7"; ws4["E5"] = "='3_Колонны'!D6"; ws4["F5"] = "='3_Колонны'!D5"
# d min
ws4["D6"] = "—"
ws4["E6"] = f"=D14+2*{DELTA_S}"
ws4["F6"] = f"=E14+2*{DELTA_S}"
# D
ws4["D7"] = f"={DPROD}"
amin(ws4, "E7", CD, Cd, "E6")
amin(ws4, "F7", CD, Cd, "F6")
for col in "DEF":
    ws4[f"{col}8"] = f"=INDEX({CN},MATCH({col}7,{CD},0))"
# s — выбор (ввод), d — по таблице
ws4["D9"] = "по расчёту на прочность"
ws4["E9"] = WALL_CHOICE["Промежуточная"]
ws4["F9"] = WALL_CHOICE["Кондуктор"]
ws4["D10"] = "—"
ws4["E10"] = f"=SUMIFS({Cd},{CD},E7,{CS},E9)"
ws4["F10"] = f"=SUMIFS({Cd},{CD},F7,{CS},F9)"
for col in "DEF":
    ws4[f"{col}11"] = f"=INDEX({CM},MATCH({col}7,{CD},0))"
    amin(ws4, f"{col}12", DVAL, DLIM, f"{col}8")
    ws4[f"{col}13"] = f"={col}11+2*{col}12"
    amin(ws4, f"{col}14", BB, BB, f"{col}13")
    ws4[f"{col}15"] = f"=({col}14-{col}11)/2"
ws4["D16"] = "—"
ws4["E16"] = "=(E10-D14)/2"
ws4["F16"] = "=(F10-E14)/2"
box(ws4, "B5:F16")
box(ws4, "B5:B16", align=LEFT)
box(ws4, "D5:F16", fmt="0.0")
box(ws4, "E9:F9", INPUT, fmt="0.0")
for rr in (7, 14):
    box(ws4, f"D{rr}:F{rr}", RESULT, bold=True, fmt="0.0")
ws4["B18"] = "Итог (отечественная методика)"; ws4["B18"].font = BOLD
ws4["B19"] = '="Эксплуатационная "&FIXED(D7,1)&" мм — долото "&FIXED(D14,1)&" мм; L = "&D5&" м"'
ws4["B20"] = '="Промежуточная "&FIXED(E7,1)&" мм — долото "&FIXED(E14,1)&" мм; L = "&E5&" м"'
ws4["B21"] = '="Кондуктор "&FIXED(F7,1)&" мм — долото "&FIXED(F14,1)&" мм; L = "&F5&" м"'
ws4["B23"] = ("Примечания: 1) муфты — ОТТМ (трапецеидальная резьба, лучше герметичность для газовой скважины); "
              "2) толщины стенок предварительные — окончательно по расчёту на прочность; для промежуточной "
              "допустимы s = 6,9…10,4 мм (d ≥ 156 мм), для кондуктора — любая из таблицы (d ≥ 232,3 мм); "
              "3) ОК 244,5 мм для кондуктора не подходит: её наибольший d = 228,7 мм < 232,3 мм.")
ws4.merge_cells("B23:F25")
ws4["B23"].alignment = LEFT
ws4["B23"].font = Font(italic=True, size=9)

# =====================================================================
# Лист 5. Зарубежная методика и сравнение
# =====================================================================
ws5 = wb.create_sheet("5_Зарубежная методика")
wb.move_sheet(ws5, offset=-2)
widths(ws5, [3, 20, 12, 12, 12, 12, 14, 14, 14, 16])
title(ws5, "B1", "Подбор по зарубежной методике (схема С01, стр. 17–18, стандартный зазор) и сравнение")
hdr = ["Колонна", "ОК, дюйм", "ОК, мм", "долото, дюйм", "долото, мм",
       "ОК РФ, мм", "долото РФ, мм", "разница долот, мм", "Δ факт = (Dдол − dм)/2"]
for j, h in enumerate(hdr):
    ws5.cell(4, 2 + j, h)
box(ws5, "B4:J4", HEAD, True)
ru_col = {"Эксплуатационная": "D", "Промежуточная": "E", "Кондуктор": "F"}
for i, (name, ci, cm, bi, bm) in enumerate(FOREIGN):
    r = 5 + i
    c = ru_col[name]
    ws5.cell(r, 2, name); ws5.cell(r, 3, ci); ws5.cell(r, 4, cm); ws5.cell(r, 5, bi); ws5.cell(r, 6, bm)
    ws5.cell(r, 7, f"='4_Диаметры (РФ)'!{c}7")
    ws5.cell(r, 8, f"='4_Диаметры (РФ)'!{c}14")
    ws5.cell(r, 9, f"=F{r}-H{r}")
    ws5.cell(r, 10, f"=(F{r}-INDEX({CM},MATCH(D{r},{CD},0)))/2")
box(ws5, "B5:J7", fmt="0.0")
box(ws5, "D5:D7", RESULT, fmt="0.0")
box(ws5, "F5:F7", RESULT, fmt="0.0")

ws5["B9"] = "Проверка зарубежного варианта по отечественным нормам"; ws5["B9"].font = BOLD
for j, h in enumerate(["Колонна", "Δ треб., мм", "Δ факт, мм", "соответствует?", "долото есть в ГОСТ 20692?"]):
    ws5.cell(10, 2 + j, h)
box(ws5, "B10:F10", HEAD, True)
for i in range(3):
    r = 11 + i
    sr = 5 + i
    ws5.cell(r, 2, f"=B{sr}")
    amin(ws5, f"C{r}", DVAL, DLIM, f"INDEX({CN},MATCH(D{sr},{CD},0))")
    ws5.cell(r, 4, f"=J{sr}")
    ws5.cell(r, 5, f'=IF(D{r}>=C{r},"да","нет (Δ меньше нормы)")')
    ws5.cell(r, 6, ["нет (ближайшие 151,0 и 161,0)", "да", "да (311,1)"][i])
box(ws5, "B11:F13", fmt="0.0")

ws5["B15"] = "Альтернатива по схеме: долото 8 3/4\" (222,3 мм) под ОК 7\" → ОК 10 3/4\" (273,1 мм) → долото 14 1/2\" (368,3 мм) станд. / 12 1/2\" (317,5 мм) малый зазор."
ws5["B16"] = "Вывод: Ø эксплуатационной и промежуточной колонн совпадают (114,3 и 177,8 мм). Зарубежная методика допускает меньший зазор под 7\" колонну (долото 215,9 вместо 222,3 мм), поэтому кондуктор получается 244,5 мм вместо 273,1 мм (долото 311,2 вместо 349,2 мм)."
for r in (15, 16):
    ws5.merge_cells(f"B{r}:J{r}")
    ws5[f"B{r}"].alignment = LEFT
    ws5.row_dimensions[r].height = 34

for r in (4, 13, 26):
    ws2.row_dimensions[r].height = 36

# оформление печати
for w in wb.worksheets:
    w.sheet_view.zoomScale = 100
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth = 1
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToHeight = 0

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
